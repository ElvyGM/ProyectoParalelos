from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "Entrega" / "Investigacion-Computacion-Paralela-Elvy-1222572.docx"
OUT.parent.mkdir(parents=True, exist_ok=True)

topics = [
    ("1. Árboles balanceados", "Un árbol balanceado mantiene alturas similares entre sus subárboles para evitar recorridos muy largos. En un AVL, la diferencia de altura entre los hijos de cada nodo se limita y se corrige con rotaciones. En árboles rojo-negro se aplican reglas de color y rotaciones que garantizan una altura logarítmica.", "Su búsqueda, inserción y eliminación se mantienen en O(log n). En sistemas paralelos ayudan a repartir búsquedas por ramas y a reducir puntos calientes. Una base de datos puede usar estructuras balanceadas para localizar registros e índices sin recorrer todos los datos.", "En UserFlow, PostgreSQL utiliza índices para acelerar búsquedas por identificador o correo. El concepto explica por qué un índice bien organizado responde mejor que un recorrido completo de usuarios."),
    ("2. Técnica Pointer Jumping", "Pointer Jumping acelera el seguimiento de enlaces haciendo que cada elemento salte hacia el elemento apuntado por su sucesor. En cada ronda, la distancia recorrida se duplica. Una cadena de longitud n puede resolverse en aproximadamente O(log n) rondas paralelas.", "Se usa en listas enlazadas, cálculo de raíces en bosques y componentes conexos. Cada procesador actualiza su puntero con información de la ronda anterior, por lo que se necesita sincronización entre etapas.", "Puede imaginarse una cadena de supervisores: en una ronda cada empleado conoce a su jefe; después conoce al jefe de su jefe. Repetir el salto permite llegar rápidamente a la raíz de la jerarquía."),
    ("3. Divide y vencerás", "Divide y vencerás separa un problema en subproblemas más pequeños, resuelve cada uno y combina sus resultados. Merge Sort, Quick Sort y la búsqueda binaria son ejemplos clásicos.", "Los subproblemas independientes pueden ejecutarse de manera concurrente. El rendimiento depende del costo de dividir, combinar y sincronizar. Si una rama recibe mucho más trabajo, se pierde parte de la ventaja paralela.", "La carga de varios archivos puede dividirse en una tarea por archivo. Future.wait inicia esas tareas y combina los resultados cuando todas terminan."),
    ("4. Particionamiento", "Particionar significa distribuir datos o trabajo en grupos. El particionamiento puede ser por bloques, por rangos, cíclico o mediante hash.", "Un buen reparto reduce comunicación y mantiene ocupados los recursos. Un reparto deficiente deja procesos ociosos mientras otros siguen trabajando. En bases de datos, particionar también puede reducir la cantidad de filas examinadas.", "Cloud Run crea varias instancias cuando aumenta la demanda. Cada solicitud se asigna a una instancia disponible. Cloud SQL y el almacenamiento central conservan el estado compartido."),
    ("5. Técnica de Pipelining", "Pipelining divide un proceso en etapas y permite que diferentes elementos avancen simultáneamente por etapas distintas. Se parece a una línea de ensamblaje.", "El rendimiento queda limitado por la etapa más lenta. Los buffers permiten absorber diferencias temporales, pero un exceso aumenta memoria y latencia. Los procesadores usan pipelines para solapar búsqueda, decodificación y ejecución de instrucciones.", "El pipeline de GitHub Actions ejecuta restauración, pruebas, compilación, construcción Docker y despliegue. Cada fase produce una salida verificable para la siguiente."),
    ("6. Aceleramiento en cascada", "El aceleramiento en cascada combina optimizaciones en varios niveles. Una mejora prepara o amplifica la siguiente, por ejemplo algoritmo eficiente, paralelismo, caché, red y escalado de infraestructura.", "Las mejoras no siempre se multiplican de forma perfecta porque aparecen límites secuenciales, comunicación y contención. La ley de Amdahl recuerda que la parte no paralelizable limita el aceleramiento total.", "UserFlow combina Future.wait en el cliente, endpoints asíncronos en ASP.NET Core, índices y conexiones de PostgreSQL, contenedores y escalado de Cloud Run."),
    ("7. Técnica Symmetry Breaking", "Cuando varios procesos tienen el mismo estado y las mismas opciones, pueden intentar la misma acción a la vez. Symmetry Breaking introduce una diferencia, como un identificador, prioridad, turno, elección aleatoria o líder.", "La técnica permite elegir coordinadores, asignar recursos y evitar decisiones idénticas. Es fundamental en algoritmos distribuidos y en problemas como elección de líder o coloreado de grafos.", "Los identificadores únicos de usuarios, archivos, ejecuciones y recursos de Terraform permiten distinguir participantes y aplicar decisiones deterministas."),
    ("8. Cluster", "Un cluster es un conjunto de computadoras que colaboran como una plataforma. Los nodos pueden repartir cómputo, almacenar réplicas o asumir el trabajo de un nodo que falla.", "Sus ventajas incluyen escalabilidad horizontal y tolerancia a fallos. Sus retos incluyen consistencia, latencia de red, observabilidad y coordinación.", "Cloud Run oculta la administración del cluster: Google inicia y detiene instancias del contenedor. El desarrollador configura límites de escalado y trabaja con el servicio lógico."),
    ("9. Balanceo de cargas", "El balanceador distribuye solicitudes entre servidores saludables. Puede usar ronda circular, menor número de conexiones, pesos, ubicación o hashing.", "El objetivo no es solo repartir por igual, sino reducir latencia y evitar sobrecarga. Las comprobaciones de salud impiden enviar tráfico a instancias defectuosas.", "API Gateway recibe una dirección pública y Cloud Run distribuye solicitudes entre instancias. El endpoint /health permite comprobar que el backend responde."),
    ("10. Sistema operativo", "El sistema operativo administra procesadores, memoria, archivos, dispositivos y seguridad. Su planificador decide qué hilo utiliza cada núcleo y durante cuánto tiempo.", "Los procesos aíslan espacios de memoria; los hilos comparten recursos del proceso. Cambiar de contexto tiene costo, por lo que crear más hilos no garantiza mayor velocidad.", "El contenedor ejecuta ASP.NET Core sobre Linux administrado por Cloud Run. Android administra el proceso Flutter, la red y el acceso a archivos del teléfono."),
    ("11. Procesador de tareas", "Un procesador de tareas recibe trabajos, los coloca en una cola y asigna trabajadores. Puede reintentar, limitar concurrencia, registrar resultados y mover fallos persistentes a una cola especial.", "Este patrón desacopla la recepción del trabajo de su ejecución. Resulta útil para imágenes, correos, reportes o tareas que no deben bloquear una solicitud web.", "Como ampliación, UserFlow podría enviar cada archivo a una cola para análisis antivirus o generación de miniaturas, mientras la API responde que el trabajo fue aceptado."),
    ("12. Nexus", "Nexus Repository es un administrador de artefactos. Guarda paquetes, bibliotecas e imágenes de contenedor con versiones y permisos. También puede funcionar como proxy y caché de repositorios externos.", "Un repositorio central mejora trazabilidad y evita depender de descargas externas en cada compilación. La política de retención evita acumular versiones innecesarias.", "El proyecto usa Artifact Registry, que cumple una función equivalente para la imagen userflow-api. Cloud Run despliega la etiqueta publicada por el pipeline."),
    ("13. Repositorios de imágenes", "Un repositorio de imágenes almacena imágenes OCI/Docker por nombre y etiqueta. Cada imagen contiene capas inmutables y puede identificarse por un digest criptográfico.", "El flujo normal consiste en construir, analizar, publicar y desplegar. Las etiquetas son cómodas, mientras que los digests ofrecen una referencia inmutable y más reproducible.", "Cloud Build construye la imagen desde el Dockerfile, la publica en Artifact Registry y Cloud Run inicia el contenedor desde ese repositorio."),
    ("14. Pub/Sub", "Pub/Sub aplica el modelo publicar-suscribir. Un productor publica mensajes en un tema sin conocer a los consumidores. Una o más suscripciones entregan los mensajes a trabajadores independientes.", "El patrón facilita sistemas desacoplados y escalables. El consumidor debe tolerar duplicados, confirmar mensajes procesados y definir reintentos o una cola de mensajes fallidos.", "UserFlow podría publicar un evento ArchivoSubido. Distintos consumidores podrían registrar auditoría, generar miniaturas o notificar al usuario sin modificar el endpoint principal."),
    ("15. GitHub Actions", "GitHub Actions automatiza tareas mediante workflows YAML activados por eventos como push o pull request. Los jobs contienen pasos que se ejecutan en runners.", "Los secretos evitan guardar credenciales en el repositorio. Workload Identity Federation mejora la seguridad porque entrega credenciales temporales y evita claves JSON permanentes.", "El workflow de UserFlow compila y prueba .NET, autentica con GCP, construye Docker, ejecuta Terraform y despliega Cloud Run. La ejecución verde constituye evidencia reproducible de CI/CD."),
]

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)
styles["Title"].font.name = "Aptos Display"
styles["Title"].font.size = Pt(28)
styles["Title"].font.color.rgb = RGBColor(39, 57, 125)
for name in ("Heading 1", "Heading 2"):
    styles[name].font.name = "Aptos Display"
    styles[name].font.color.rgb = RGBColor(39, 57, 125)

p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Técnicas de computación paralela y plataformas cloud")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Investigación y actividad de exposición")
r.bold = True
r.font.size = Pt(16)
for line in ["ELVY ANTONIO GRULLÓN MOREL", "Matrícula: 1222572", "Caso práctico: UserFlow", "Octubre de 2026"]:
    p = doc.add_paragraph(line)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_page_break()

doc.add_heading("Introducción", level=1)
doc.add_paragraph("La computación paralela busca ejecutar trabajo simultáneamente para reducir tiempo, aumentar capacidad o mantener disponible un sistema. La concurrencia organiza tareas que progresan durante el mismo intervalo; el paralelismo ejecuta trabajo al mismo tiempo en varios núcleos o máquinas. Esta investigación explica técnicas algorítmicas, componentes de infraestructura y herramientas de automatización, y las conecta con la aplicación UserFlow.")
doc.add_heading("Objetivos", level=1)
for item in ["Definir cada concepto con lenguaje claro.", "Explicar su relación con concurrencia, paralelismo o sistemas distribuidos.", "Identificar beneficios, límites y ejemplos.", "Relacionar los conceptos con Flutter, ASP.NET Core y Google Cloud."]:
    doc.add_paragraph(item, style="List Bullet")

for title, definition, importance, example in topics:
    doc.add_heading(title, level=1)
    doc.add_heading("Concepto", level=2)
    doc.add_paragraph(definition)
    doc.add_heading("Importancia y funcionamiento", level=2)
    doc.add_paragraph(importance)
    doc.add_heading("Aplicación o ejemplo", level=2)
    doc.add_paragraph(example)

doc.add_heading("Actividad práctica", level=1)
doc.add_paragraph("Demostración propuesta: comparar la carga secuencial y concurrente de varios archivos. En el modo secuencial se espera cada respuesta antes de iniciar la siguiente. En el modo concurrente se crean todos los Future y Future.wait espera el grupo. La pantalla registra milisegundos y calcula el porcentaje de mejora. El resultado depende de red, tamaño de archivo, límites del servidor y número de conexiones; por eso debe repetirse la prueba y explicar que concurrencia no garantiza siempre una aceleración.")

table = doc.add_table(rows=1, cols=3)
table.style = "Table Grid"
hdr = table.rows[0].cells
hdr[0].text, hdr[1].text, hdr[2].text = "Parte", "Qué mostrar", "Qué explicar"
rows = [
    ("Future.wait", "Código del ViewModel", "Inicia futuros independientes y espera todos los resultados."),
    ("Pipeline", "Ejecución verde en Actions", "Automatiza build, test, Docker, Terraform y despliegue."),
    ("Cluster/balanceo", "Cloud Run", "Google distribuye solicitudes y escala instancias."),
    ("Repositorio", "Artifact Registry", "Conserva la imagen Docker que se despliega."),
    ("Seguridad", "JWT y Secret Manager", "El token identifica al usuario y los secretos no están en Git."),
]
for values in rows:
    cells = table.add_row().cells
    for i, value in enumerate(values):
        cells[i].text = value

doc.add_heading("Conclusiones", level=1)
doc.add_paragraph("Las técnicas estudiadas resuelven problemas distintos. Divide y vencerás, particionamiento y Pointer Jumping organizan algoritmos; pipelining y aceleramiento en cascada aprovechan etapas y niveles; Symmetry Breaking coordina participantes; clusters, balanceadores, sistemas operativos y procesadores de tareas administran recursos; Nexus, repositorios de imágenes, Pub/Sub y GitHub Actions permiten construir y operar plataformas distribuidas. UserFlow integra varios de estos principios en una solución demostrable.")

doc.add_heading("Referencias", level=1)
refs = [
    "Cormen, T. et al. Introduction to Algorithms. MIT Press.",
    "Grama, A. et al. Introduction to Parallel Computing. Addison-Wesley.",
    "Tanenbaum, A. Modern Operating Systems. Pearson.",
    "Google Cloud. Documentación de Cloud Run, API Gateway, Cloud SQL, Pub/Sub y Artifact Registry. https://cloud.google.com/docs",
    "GitHub Docs. Understanding GitHub Actions. https://docs.github.com/actions",
    "Sonatype Help. Nexus Repository documentation. https://help.sonatype.com/en/sonatype-nexus-repository.html",
    "Flutter API. Future.wait. https://api.dart.dev/dart-async/Future/wait.html",
]
for ref in refs:
    doc.add_paragraph(ref, style="List Number")

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run("ELVY ANTONIO GRULLÓN MOREL · 1222572")
run.font.size = Pt(8)
run.font.color.rgb = RGBColor(100, 100, 100)

doc.core_properties.title = "Técnicas de computación paralela y plataformas cloud"
doc.core_properties.author = "ELVY ANTONIO GRULLÓN MOREL"
doc.core_properties.subject = "Investigación y exposición"
doc.save(OUT)
print(OUT)
